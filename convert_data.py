import openpyxl
import json
import re
import os

def parse_date(v):
    if v is None:
        return None
    s = str(v).split('.')[0].strip()
    if len(s) == 6:
        # e.g. 220831 -> 2022-08-31
        return f"20{s[0:2]}-{s[2:4]}-{s[4:6]}"
    elif len(s) == 8:
        return f"{s[0:4]}-{s[4:6]}-{s[6:8]}"
    return s

def clean_note_text(item, notes):
    return " ".join(notes)

excel_path = "家裡、老爸的照顧花費＆照顧紀錄 的副本.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)

# 1. 就醫、照顧花費
s1 = wb["就醫、照顧花費記帳"]
s1_records = []
s1_id = 1
for r in range(4, s1.max_row + 1):
    d = parse_date(s1.cell(row=r, column=2).value)
    if not d:
        continue
    raw_item = str(s1.cell(row=r, column=3).value or "").strip()
    amt = s1.cell(row=r, column=5).value
    note = s1.cell(row=r, column=6).value or ""
    mom = s1.cell(row=r, column=7).value

    amount = float(amt) if amt is not None else 0.0
    notes = [str(note).strip()] if str(note).strip() else []

    if mom is not None:
        try:
            mom_num = float(mom)
            amount += mom_num
            notes.append("老媽")
        except ValueError:
            notes.append(f"老媽: {mom}")

    # Map item to: 看診、長照費、住院費、看護費、其他
    if any(k in raw_item for k in ["看診", "門診", "回診", "急診"]):
        item_cat = "看診"
        if raw_item != "看診" and raw_item not in clean_note_text(raw_item, notes):
            notes.insert(0, f"原項目: {raw_item}")
    elif any(k in raw_item for k in ["長照"]):
        item_cat = "長照費"
        if raw_item != "長照費" and raw_item not in clean_note_text(raw_item, notes):
            notes.insert(0, f"原項目: {raw_item}")
    elif any(k in raw_item for k in ["住院", "急診預繳"]):
        item_cat = "住院費"
        if raw_item != "住院費" and raw_item not in clean_note_text(raw_item, notes):
            notes.insert(0, f"原項目: {raw_item}")
    elif any(k in raw_item for k in ["看護"]):
        item_cat = "看護費"
        if raw_item != "看護費" and raw_item not in clean_note_text(raw_item, notes):
            notes.insert(0, f"原項目: {raw_item}")
    else:
        item_cat = "其他"
        if raw_item and raw_item != "其他":
            notes.insert(0, f"原項目: {raw_item}")

    final_note = "; ".join(notes)
    s1_records.append({
        "id": f"med_{s1_id}",
        "date": d,
        "item": item_cat,
        "amount": int(amount) if amount.is_integer() else amount,
        "note": final_note
    })
    s1_id += 1

# 2. 家用花費、繳款紀錄
s2 = wb["家用花費、繳款紀錄"]
s2_records = []
s2_id = 1
for r in range(4, s2.max_row + 1):
    d = parse_date(s2.cell(row=r, column=2).value)
    if not d:
        continue
    raw_item = str(s2.cell(row=r, column=3).value or "").strip()
    amt = s2.cell(row=r, column=6).value
    note = s2.cell(row=r, column=7).value or ""
    amount = float(amt) if amt is not None else 0.0
    notes = [str(note).strip()] if str(note).strip() else []

    # Map item to: 台電、中華電信、瓦斯、北水、房貸轉帳、墓園管理費、其他
    if "台電" in raw_item:
        item_cat = "台電"
    elif "中華電信" in raw_item:
        item_cat = "中華電信"
    elif "瓦斯" in raw_item:
        item_cat = "瓦斯"
    elif "北水" in raw_item or "水" in raw_item:
        item_cat = "北水"
    elif "轉帳" in raw_item:
        item_cat = "房貸轉帳"
    elif "墓園管理費" in raw_item:
        item_cat = "墓園管理費"
        if "土城" in raw_item:
            notes.insert(0, "土城")
        elif "八里" in raw_item:
            notes.insert(0, "八里")
    else:
        item_cat = "其他"
        if raw_item and raw_item != "其他":
            notes.insert(0, f"原項目: {raw_item}")

    final_note = "; ".join(notes)
    s2_records.append({
        "id": f"house_{s2_id}",
        "date": d,
        "item": item_cat,
        "amount": int(amount) if amount.is_integer() else amount,
        "note": final_note
    })
    s2_id += 1

# 3. 老哥轉帳記錄
s4 = wb["老哥轉帳記錄"]
s4_records = []
s4_id = 1
for r in range(5, s4.max_row + 1):
    d = parse_date(s4.cell(row=r, column=2).value)
    if not d:
        continue
    twd = s4.cell(row=r, column=7).value
    note = s4.cell(row=r, column=8).value or ""
    if twd is not None and twd != "":
        try:
            amt = float(twd)
            s4_records.append({
                "id": f"transfer_{s4_id}",
                "date": d,
                "amount": int(amt) if amt.is_integer() else amt,
                "note": str(note).strip()
            })
            s4_id += 1
        except ValueError:
            pass

data = {
    "medicalCare": s1_records,
    "household": s2_records,
    "transfers": s4_records
}

with open("data.json", "w", encoding="utf-8") as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

with open("data.js", "w", encoding="utf-8") as f:
    f.write("// 預設初始歷史資料集 (由 Excel 萃取匯入)\n")
    f.write("window.INITIAL_DATA = ")
    json.dump(data, f, ensure_ascii=False, indent=2)
    f.write(";\n")

print(f"Medical: {len(s1_records)}, sum={sum(x['amount'] for x in s1_records)}")
print(f"Household: {len(s2_records)}, sum={sum(x['amount'] for x in s2_records)}")
print(f"Transfers: {len(s4_records)}, sum={sum(x['amount'] for x in s4_records)}")
